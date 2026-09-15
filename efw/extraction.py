"""Row-by-row extraction of Finding objects from a worksheet, value
validation (Risk Level / Impact / Likelihood - warns but never raises),
and the section-grouping logic that ties it all together.
"""

from typing import Optional

from openpyxl.worksheet.worksheet import Worksheet

from .constants import (
    RISK_LEVELS,
    IMPACT_LEVELS,
    LIKELIHOOD_LEVELS,
    VERIFICATION_CANONICAL_LABELS,
)
from .diagnostics import warn
from .excel_source import section_title_for_row
from .models import Finding, HeaderMap, SAFinding, SAHeaderMap
from .text_utils import (
    clean_text,
    extract_date,
    normalize,
    normalize_verification_status,
    split_paragraphs,
)

# =============================================================================
# Step 3-4: extract + validate findings
# =============================================================================


def get_val(ws: Worksheet, row: int, col: Optional[int]) -> str:
    if col is None:
        return ""
    return clean_text(ws.cell(row=row, column=col).value)


def validate_risk_level(value: str, finding_id: str, row: int) -> None:
    if value and normalize(value) not in RISK_LEVELS:
        warn(
            f'Row {row} (Finding "{finding_id}"): unexpected Risk Level '
            f'"{value}". Expected one of {sorted(RISK_LEVELS)}.'
        )


def validate_impact(value: str, finding_id: str, row: int) -> None:
    if value and normalize(value) not in IMPACT_LEVELS:
        warn(
            f'Row {row} (Finding "{finding_id}"): unexpected Impact '
            f'"{value}". Expected one of {sorted(IMPACT_LEVELS)}.'
        )


def validate_likelihood(value: str, finding_id: str, row: int) -> None:
    if value and normalize(value) not in LIKELIHOOD_LEVELS:
        warn(
            f'Row {row} (Finding "{finding_id}"): unexpected Likelihood '
            f'"{value}". Expected one of {sorted(LIKELIHOOD_LEVELS)}.'
        )


def extract_findings(
    ws: Worksheet, hmap: HeaderMap, header_row: int, min_col: int, max_col: int, max_row: int
) -> list[tuple[Optional[str], list[Finding]]]:
    """Returns a list of (section_title_or_None, [Finding, ...]) groups, in
    top-to-bottom order, matching the layout of the source sheet."""

    groups: list[tuple[Optional[str], list[Finding]]] = []
    current_section: Optional[str] = None
    current_findings: list[Finding] = []

    for row in range(header_row + 1, max_row + 1):
        section = section_title_for_row(ws, row, min_col, max_col)
        if section is not None:
            if current_findings or current_section is not None:
                groups.append((current_section, current_findings))
            current_section = section
            current_findings = []
            continue

        finding_id = get_val(ws, row, hmap.id_col)
        finding_title = get_val(ws, row, hmap.finding_col)

        # Skip fully blank rows
        row_values = [
            get_val(ws, row, c)
            for c in (
                hmap.id_col,
                hmap.finding_col,
                hmap.affected_col,
                hmap.risk_level_col,
                hmap.risk_description_col,
            )
        ]
        if not any(row_values):
            continue

        risk_level = get_val(ws, row, hmap.risk_level_col)
        impact = get_val(ws, row, hmap.impact_col)
        likelihood = get_val(ws, row, hmap.likelihood_col)
        affected = get_val(ws, row, hmap.affected_col)
        risk_description = split_paragraphs(ws.cell(row=row, column=hmap.risk_description_col).value) if hmap.risk_description_col else []
        recommended_safeguards = split_paragraphs(ws.cell(row=row, column=hmap.recommended_safeguards_col).value) if hmap.recommended_safeguards_col else []
        owasp_top_10 = split_paragraphs(ws.cell(row=row, column=hmap.owasp_top_10_col).value) if hmap.owasp_top_10_col else []

        # Verification: take the LAST non-empty column, left-to-right.
        verification_raw = ""
        verification_header = ""
        verification_col_used: Optional[int] = None
        for col, header_text in hmap.verification_cols:
            val = get_val(ws, row, col)
            if val:
                verification_raw = val
                verification_header = header_text
                verification_col_used = col

        # Normalize to one of the 5 canonical labels (prefix match), dropping
        # any trailing explanatory text. Warn if nothing matches.
        verification_status, matched = normalize_verification_status(verification_raw)
        if not matched:
            warn(
                f'Row {row} (Finding "{finding_id}"): unexpected verification '
                f'status "{verification_raw}". Expected it to start with one '
                f"of {VERIFICATION_CANONICAL_LABELS} (case-insensitive)."
            )

        date_str = extract_date(verification_header) if verification_header else None
        if date_str:
            verification_date_label = f"Rectification status as of {date_str}"
        elif verification_header:
            verification_date_label = f"Rectification status ({verification_header})"
        else:
            verification_date_label = "Rectification status as of dd MMM yyyy"

        validate_risk_level(risk_level, finding_id, row)
        validate_impact(impact, finding_id, row)
        validate_likelihood(likelihood, finding_id, row)

        current_findings.append(
            Finding(
                row=row,
                finding_id=finding_id,
                finding_title=finding_title,
                affected=affected,
                risk_level=risk_level,
                impact=impact,
                likelihood=likelihood,
                risk_description=risk_description,
                recommended_safeguards=recommended_safeguards,
                owasp_top_10=owasp_top_10,
                verification_status=verification_status,
                verification_date_label=verification_date_label,
                verification_col_used=verification_col_used,
            )
        )

    if current_findings or current_section is not None:
        groups.append((current_section, current_findings))

    return groups


def group_verification_status_label(findings, hmap) -> str:
    """Derive a SINGLE 'Rectification Status as of <date>' column header for
    an entire section's (or, for SA, the WHOLE table's) summary table (used
    by the "landscape-detail" and "sa-detail" formats, where the
    rectification/verification date is a per-TABLE column header rather
    than a per-ROW label as in "portrait-detail").

    Convention (reusing the same date-extraction logic as the per-row
    label): among all verification columns actually used by ANY finding/item
    in `findings`, pick the RIGHTMOST one (i.e. the latest verification
    round that applies) and extract its date the same way the per-row
    label does. Falls back to a generic placeholder if no verification data
    is available.

    Works with EITHER `Finding`+`HeaderMap` (SRA) or `SAFinding`+`SAHeaderMap`
    (SA) objects, since both expose the same `.verification_cols` /
    `.verification_col_used` attribute names.
    """
    if not hmap.verification_cols:
        return "Rectification Status as of dd MMM yyyy"

    used_cols = {f.verification_col_used for f in findings if f.verification_col_used is not None}
    if not used_cols:
        return "Rectification Status as of dd MMM yyyy"

    max_col = max(used_cols)
    header_text = next((h for c, h in hmap.verification_cols if c == max_col), "")

    date_str = extract_date(header_text) if header_text else None
    if date_str:
        return f"Rectification Status as of {date_str}"
    elif header_text:
        return f"Rectification Status ({header_text})"
    return "Rectification Status as of dd MMM yyyy"


# =============================================================================
# SA (Security Audit) extraction - "SA Follow-up" table
# =============================================================================


def extract_sa_findings(
    ws: Worksheet, hmap: SAHeaderMap, header_row: int, min_col: int, max_col: int, max_row: int
) -> list[SAFinding]:
    """Returns a FLAT list of SAFinding, in top-to-bottom order. Unlike the
    SRA table, the SA table has no section-header rows in its spec, so no
    grouping is performed."""
    items: list[SAFinding] = []

    for row in range(header_row + 1, max_row + 1):
        item_id = get_val(ws, row, hmap.id_col)

        row_values = [
            get_val(ws, row, c)
            for c in (hmap.id_col, hmap.items_col, hmap.affected_col, hmap.findings_col)
        ]
        if not any(row_values):
            continue

        items_to_check = split_paragraphs(ws.cell(row=row, column=hmap.items_col).value) if hmap.items_col else []
        affected = get_val(ws, row, hmap.affected_col)
        findings = split_paragraphs(ws.cell(row=row, column=hmap.findings_col).value) if hmap.findings_col else []
        recommended_safeguards = split_paragraphs(ws.cell(row=row, column=hmap.recommended_safeguards_col).value) if hmap.recommended_safeguards_col else []

        # Verification: take the LAST non-empty column, left-to-right (same
        # convention as SRA - handles non-contiguous verification columns
        # naturally, since hmap.verification_cols already lists them in
        # left-to-right column order regardless of what's interleaved
        # between them).
        verification_raw = ""
        verification_header = ""
        verification_col_used: Optional[int] = None
        for col, header_text in hmap.verification_cols:
            val = get_val(ws, row, col)
            if val:
                verification_raw = val
                verification_header = header_text
                verification_col_used = col

        verification_status, matched = normalize_verification_status(verification_raw)
        if not matched:
            warn(
                f'Row {row} (Item "{item_id}"): unexpected verification '
                f'status "{verification_raw}". Expected it to start with one '
                f"of {VERIFICATION_CANONICAL_LABELS} (case-insensitive)."
            )

        date_str = extract_date(verification_header) if verification_header else None
        if date_str:
            verification_date_label = f"Rectification status as of {date_str}"
        elif verification_header:
            verification_date_label = f"Rectification status ({verification_header})"
        else:
            verification_date_label = "Rectification status as of dd MMM yyyy"

        items.append(
            SAFinding(
                row=row,
                item_id=item_id,
                items_to_check=items_to_check,
                affected=affected,
                findings=findings,
                recommended_safeguards=recommended_safeguards,
                verification_status=verification_status,
                verification_date_label=verification_date_label,
                verification_col_used=verification_col_used,
            )
        )

    return items
