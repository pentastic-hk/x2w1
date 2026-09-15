"""The "landscape-detail" output format: A4 landscape, one summary
Word table PER SECTION (one row per finding).
"""

from typing import Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Cm, Mm, Pt

from ..constants import (
    DEFAULT_SECTION_NUMBER,
    FONT_NAME,
    FONT_SIZE,
    HEADER_ROW_FILL,
    HEADER_ROW_FONT_COLOR,
    RISK_LEVEL_COLORS,
)
from ..diagnostics import warn
from ..docx_utils import (
    _apply_base_styles,
    _set_cell_fill,
    _set_cell_text,
    _set_repeat_header_row,
    _set_table_column_widths,
)
from ..extraction import group_verification_status_label
from ..models import Finding, HeaderMap
from ..text_utils import normalize

# =============================================================================
# "landscape-detail" format: one summary table per section, A4 landscape
# =============================================================================

# Content width targeted for A4 landscape with the margins set in
# _setup_a4_landscape() below: 29.7cm - (1.5cm left + 1.5cm right) = 26.7cm.
#
# Each entry is (key, header_label_or_None, width). `key` is used to look up
# the corresponding value for each finding (see _landscape_row_values()).
# The "OWASP Top 10" column is OPTIONAL: it is only included when the source
# workbook actually has an "OWASP Top 10" column (hmap.owasp_top_10_col is
# not None) - decided ONCE per document/workbook (not per-section), so every
# section's table has the same shape.
#
# Two separate width sets are defined (rather than simply carving the new
# column's width out of a single existing column) so that adding "OWASP Top
# 10" doesn't make any one column uncomfortably narrow - the extra width is
# instead spread proportionally across several columns. Both sets sum to
# the same 26.7cm total content width.
_LANDSCAPE_WIDTHS_NO_OWASP = {
    "index": Cm(1.2),
    "findings": Cm(3.5),
    "affected": Cm(3.6),
    "risk_description": Cm(5.3),
    "risk_level": Cm(2.0),
    "impact_likelihood": Cm(2.5),
    "recommendation": Cm(5.1),
    "rectification": Cm(3.5),
}
_LANDSCAPE_WIDTHS_WITH_OWASP = {
    "index": Cm(1.0),
    "findings": Cm(3.2),
    "affected": Cm(3.2),
    "risk_description": Cm(4.6),
    "risk_level": Cm(1.8),
    "impact_likelihood": Cm(2.3),
    "owasp_top_10": Cm(3.0),
    "recommendation": Cm(4.3),
    "rectification": Cm(3.3),
}


def landscape_columns(has_owasp: bool) -> list[tuple[str, Optional[str], object]]:
    """Build the ordered list of (key, header_label_or_None, width) columns
    for the "landscape-detail" summary table, including "OWASP Top 10" only
    when `has_owasp` is True (i.e. the source workbook has that optional
    column). The last column's header is dynamic ("Rectification Status as
    of <date>") and is filled in per-section at render time - its label is
    left as None here."""
    widths = _LANDSCAPE_WIDTHS_WITH_OWASP if has_owasp else _LANDSCAPE_WIDTHS_NO_OWASP
    columns: list[tuple[str, Optional[str], object]] = [
        ("index", "#", widths["index"]),
        ("findings", "Findings", widths["findings"]),
        ("affected", "Affected", widths["affected"]),
        ("risk_description", "Risk Description", widths["risk_description"]),
        ("risk_level", "Risk Level", widths["risk_level"]),
        ("impact_likelihood", "Impact / Likelihood", widths["impact_likelihood"]),
    ]
    if has_owasp:
        columns.append(("owasp_top_10", "OWASP Top 10", widths["owasp_top_10"]))
    columns.append(("recommendation", "Recommendation", widths["recommendation"]))
    columns.append(("rectification", None, widths["rectification"]))
    return columns


def _landscape_row_values(finding: Finding) -> dict[str, list[str]]:
    """Map each landscape column `key` to the paragraph(s) to render for a
    given finding. Only keys actually present in `landscape_columns(...)`
    are looked up by the caller, so it's fine for this dict to always
    include "owasp_top_10" regardless of has_owasp."""
    return {
        "index": [finding.finding_id],
        "findings": [finding.finding_title],
        "affected": _lines_from_text(finding.affected),
        "risk_description": finding.risk_description or [""],
        "risk_level": [finding.risk_level],
        "impact_likelihood": [f"{finding.impact} / {finding.likelihood}" if (finding.impact or finding.likelihood) else ""],
        "owasp_top_10": finding.owasp_top_10 or [""],
        "recommendation": finding.recommended_safeguards or [""],
        "rectification": [finding.verification_status],
    }


def _setup_a4_landscape(document: Document) -> None:
    """Configure the document's first section as A4, landscape orientation."""
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)


def _lines_from_text(text: str) -> list[str]:
    """Split an already-cleaned string (e.g. Finding.affected, which may
    retain literal '\\n' line breaks copied from the source Excel cell,
    such as multiple IP addresses) into a list of non-empty lines."""
    if not text:
        return [""]
    parts = [p.strip() for p in text.replace("\r\n", "\n").split("\n")]
    parts = [p for p in parts if p]
    return parts or [""]


def add_section_summary_table(
    document: Document,
    findings: list[Finding],
    hmap: HeaderMap,
    columns: list[tuple[str, Optional[str], object]],
) -> None:
    """Render ONE section's findings as a single summary table (one row per
    finding), per the "landscape-detail" layout. `columns` is the
    (key, header_label_or_None, width) list produced by
    landscape_columns(has_owasp) - passed in (rather than recomputed here)
    so every section in the same document uses an IDENTICAL column set."""
    n_cols = len(columns)
    column_widths = [width for _, _, width in columns]
    table = document.add_table(rows=1, cols=n_cols)
    table.style = "Table Grid"
    table.autofit = False
    _set_table_column_widths(table, column_widths)

    rectification_header = group_verification_status_label(findings, hmap)

    # ---- Header row ----
    header_row = table.rows[0]
    for i, (key, label, width) in enumerate(columns):
        cell = header_row.cells[i]
        cell.width = width
        header_text = label if label is not None else rectification_header
        _set_cell_text(cell, [header_text], bold=True, font_color=HEADER_ROW_FONT_COLOR)
        _set_cell_fill(cell, HEADER_ROW_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    _set_repeat_header_row(header_row)

    risk_level_col_index = next(i for i, (key, _, _) in enumerate(columns) if key == "risk_level")

    # ---- Data rows (one per finding) ----
    for finding in findings:
        row = table.add_row()
        cells = row.cells
        for i, (_, _, width) in enumerate(columns):
            cells[i].width = width

        row_values = _landscape_row_values(finding)

        for i, (key, _, _) in enumerate(columns):
            cell = cells[i]
            paragraphs = row_values[key]
            _set_cell_text(cell, paragraphs, bold=False)
            if i == risk_level_col_index:
                # Reuse the EXACT SAME risk-level color scheme/logic as the
                # portrait format (RISK_LEVEL_COLORS + _set_cell_fill).
                risk_hex = RISK_LEVEL_COLORS.get(normalize(finding.risk_level))
                _set_cell_fill(cell, risk_hex)
            else:
                _set_cell_fill(cell, None)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def build_landscape_document(
    groups: list[tuple[Optional[str], list[Finding]]],
    hmap: HeaderMap,
    title: str,
    section_base_number: str = DEFAULT_SECTION_NUMBER,
) -> Document:
    """Build the "landscape-detail" output: A4 landscape, one summary table
    PER SECTION (e.g. "9.1 General Control Review"), reusing the same
    Finding/HeaderMap data produced by extract_findings()."""
    document = Document()
    _apply_base_styles(document)
    _setup_a4_landscape(document)

    document.add_heading(title, level=1)

    # Decided ONCE for the whole document (not per-section), so every
    # section's table has an identical column set: include "OWASP Top 10"
    # only if that optional column exists in the source workbook.
    has_owasp = hmap.owasp_top_10_col is not None
    columns = landscape_columns(has_owasp)

    total = 0
    subsection_index = 0
    for section_title, findings in groups:
        if not findings:
            continue

        subsection_index += 1
        heading_text = f"{section_base_number}.{subsection_index} {section_title or 'Findings'}"
        document.add_heading(heading_text, level=2)

        intro = document.add_paragraph("The following issues are identified:")
        intro.paragraph_format.space_before = Pt(0)
        intro.paragraph_format.space_after = Pt(6)
        for run in intro.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(FONT_SIZE)

        add_section_summary_table(document, findings, hmap, columns)
        total += len(findings)

        # Spacer between this section's table and the next section heading.
        spacer = document.add_paragraph("")
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(0)

    if total == 0:
        warn("No findings were extracted - the output document will be empty of tables.")

    return document
