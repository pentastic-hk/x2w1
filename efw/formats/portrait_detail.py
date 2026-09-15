"""The "portrait-detail" output format: A4 portrait, one detailed
2-column Word table PER FINDING.
"""

from typing import Optional

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.shared import Cm

from ..constants import HEADER_ROW_FILL, HEADER_ROW_FONT_COLOR, RISK_LEVEL_COLORS
from ..diagnostics import warn
from ..docx_utils import (
    _add_blank_separator_paragraph,
    _apply_base_styles,
    _set_cell_fill,
    _set_cell_text,
    _set_table_column_widths,
)
from ..models import Finding
from ..text_utils import normalize

# =============================================================================
# "portrait-detail" format: one detailed table per finding, A4 portrait
# =============================================================================

LABEL_COL_WIDTH = Cm(4.2)
VALUE_COL_WIDTH = Cm(12.8)


def add_finding_table(document: Document, finding: Finding) -> None:
    table = document.add_table(rows=10, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    _set_table_column_widths(table, [LABEL_COL_WIDTH, VALUE_COL_WIDTH])

    rows_content = [
        (finding.finding_id, [finding.finding_title]),
        ("Risk Description", finding.risk_description),
        ("Risk Level", [finding.risk_level]),
        ("Impact/Likelihood", [f"{finding.impact} / {finding.likelihood}" if (finding.impact or finding.likelihood) else ""]),
        ("OWASP Top 10", finding.owasp_top_10),
        ("Affected Asset", [finding.affected]),
        ("Evidence for the finding", [""]),
        ("Recommended Safeguards", finding.recommended_safeguards),
        ("Evidence for the remedial actions", [""]),
        (finding.verification_date_label, [finding.verification_status]),
    ]

    RISK_LEVEL_ROW_INDEX = 2

    for i, (label, value_paragraphs) in enumerate(rows_content):
        row = table.rows[i]
        label_cell, value_cell = row.cells[0], row.cells[1]

        if i == 0:
            # First row: standard blue shading, white bold font (both cells)
            _set_cell_text(label_cell, [label], bold=True, font_color=HEADER_ROW_FONT_COLOR)
            _set_cell_text(value_cell, value_paragraphs, bold=True, font_color=HEADER_ROW_FONT_COLOR)
            _set_cell_fill(label_cell, HEADER_ROW_FILL)
            _set_cell_fill(value_cell, HEADER_ROW_FILL)
        else:
            _set_cell_text(label_cell, [label], bold=False)
            _set_cell_text(value_cell, value_paragraphs, bold=False)
            _set_cell_fill(label_cell, None)  # no fill

            if i == RISK_LEVEL_ROW_INDEX:
                risk_hex = RISK_LEVEL_COLORS.get(normalize(finding.risk_level))
                _set_cell_fill(value_cell, risk_hex)  # None if unrecognized -> no fill
            else:
                _set_cell_fill(value_cell, None)  # no fill

        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def build_document(groups: list[tuple[Optional[str], list[Finding]]], title: str) -> Document:
    """Build the "portrait-detail" output: A4 portrait, one 2-column Word
    table PER FINDING."""
    document = Document()
    _apply_base_styles(document)

    document.add_heading(title, level=1)

    total = 0
    for section_title, findings in groups:
        if not findings:
            continue
        if section_title:
            document.add_heading(section_title, level=2)
        for finding in findings:
            add_finding_table(document, finding)
            total += 1
            # Separate each table with 2 empty lines (no extra paragraph
            # spacing added below them).
            _add_blank_separator_paragraph(document)
            _add_blank_separator_paragraph(document)

    if total == 0:
        warn("No findings were extracted - the output document will be empty of tables.")

    return document
