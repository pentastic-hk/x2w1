"""Low-level, reusable python-docx / OOXML helper functions shared by
every output format: table column widths, cell shading, row heights,
cell text/formatting, repeating header rows, and base document styling
(font/size/color enforcement that overrides Word's theme defaults).
"""

from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .constants import FONT_NAME, FONT_SIZE

# =============================================================================
# Shared Word-building helpers (used by ALL formats)
# =============================================================================


def _set_table_column_widths(table, widths: list) -> None:
    """Explicitly (re)write the table's <w:tblGrid> to match the given
    column widths (a list of python-docx Length objects, e.g. Cm(1.2)), and
    also set <w:tblW> to the sum of those widths.

    Why this is necessary: python-docx's `cell.width = ...` only sets the
    per-cell <w:tcW>. Some renderers use <w:tblGrid> (not the individual
    per-cell widths) to lay out a table's columns when the two disagree -
    which they do by default, since add_table() always creates a tblGrid
    with the page's default width divided evenly across all columns. If we
    only set cell.width without also correcting tblGrid, columns can render
    at the wrong (evenly-distributed) widths despite tcW being "correct" in
    the underlying XML. This function keeps both in sync so the intended
    widths are honored consistently.
    """
    tbl = table._tbl
    tblGrid = tbl.tblGrid
    for gc in list(tblGrid):
        tblGrid.remove(gc)
    for w in widths:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(w.twips))
        tblGrid.append(gridCol)

    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(w.twips for w in widths)))

    # Keep every existing row's cells in sync with the corrected grid too.
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def _set_cell_fill(cell, hex_color: Optional[str]) -> None:
    """Set the cell background fill. Pass hex_color=None for 'no fill'
    (transparent / default white background)."""
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove any existing shading element first
    for existing in tcPr.findall(qn("w:shd")):
        tcPr.remove(existing)

    if hex_color is None:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "auto")
        tcPr.append(shd)
    else:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)


def _set_row_height(row, height, rule: str = "atLeast") -> None:
    """Set an explicit row height (a python-docx Length object, e.g.
    Cm(0.9)) on a table row via OOXML <w:trHeight>.

    `rule` controls how Word treats the value:
        - "atLeast" (default): the row is AT LEAST this tall, but will
          still grow taller if the cell content needs more space (safest
          choice - guarantees the minimum height without ever clipping
          text).
        - "exact": the row is forced to EXACTLY this height, even if the
          content would otherwise need more room (can clip/overlap text
          if the content doesn't fit).
    """
    trPr = row._tr.get_or_add_trPr()
    for existing in trPr.findall(qn("w:trHeight")):
        trPr.remove(existing)
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(height.twips))
    trHeight.set(qn("w:hRule"), rule)
    trPr.append(trHeight)


def _set_cell_text(
    cell,
    paragraphs: list[str],
    bold: bool = False,
    font_color: Optional[RGBColor] = None,
    alignment: Optional[WD_ALIGN_PARAGRAPH] = None,
) -> None:
    """Write one or more paragraphs of text into a table cell, applying the
    document-wide font (Times New Roman, 12pt) with no extra spacing
    between wrapped paragraphs."""
    if not paragraphs:
        paragraphs = [""]
    cell.text = ""  # clear default empty paragraph's placeholder run
    first = True
    for text in paragraphs:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(FONT_SIZE)
        run.font.bold = bold
        if font_color is not None:
            run.font.color.rgb = font_color
        if alignment is not None:
            p.paragraph_format.alignment = alignment
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)


def _add_blank_separator_paragraph(document: Document):
    """Add a blank paragraph used purely as vertical whitespace between
    tables, with all paragraph spacing zeroed out so it doesn't introduce
    any *extra* space beyond the blank line itself."""
    p = document.add_paragraph("")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        run.font.name = FONT_NAME
        run.font.size = Pt(FONT_SIZE)
    return p


def _set_repeat_header_row(row) -> None:
    """Mark a table row as a repeating header row (OOXML <w:tblHeader/>),
    so it repeats at the top of every page the table spans."""
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def _hard_set_style_font(style, name: str = FONT_NAME, size_pt: int = FONT_SIZE) -> None:
    """Force a paragraph style's run properties (rPr) to an EXPLICIT font
    name/size/color, fully overriding Word's built-in theme-based styling.

    This is needed because Word's default "Heading 1"/"Heading 2" styles
    (and similar) don't just set a plain font name/size/color - they
    reference the document THEME instead:
        <w:rFonts w:asciiTheme="majorHAnsi" w:eastAsiaTheme="majorEastAsia"
                  w:hAnsiTheme="majorHAnsi" w:cstheme="majorBidi"/>
        <w:color w:val="365F91" w:themeColor="accent1" w:themeShade="BF"/>

    Simply setting `style.font.name = ...` (as python-docx's high-level API
    does) only ADDS explicit w:ascii/w:hAnsi attributes - it does NOT
    remove the w:asciiTheme/w:hAnsiTheme/etc. attributes already present.
    When both explicit and theme font references coexist, some renderers
    (observed with LibreOffice) still preferred the THEME font over the
    explicit one, so headings kept rendering in the theme's sans-serif
    heading font (and theme accent color) instead of Times New Roman
    black/auto. This function removes ALL theme references and writes
    fully explicit values so there's no ambiguity for any renderer.
    """
    # Set the size color etc via the standard python-docx API first (keeps
    # everything consistent / handles version differences gracefully).
    style.font.name = name
    style.font.size = Pt(size_pt)

    rPr = style.element.get_or_add_rPr()

    # ---- Fonts: remove theme references, set every font slot explicitly ----
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rFonts.get(qn(theme_attr)) is not None:
            del rFonts.attrib[qn(theme_attr)]
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)

    # ---- Size: also fix complex-script size (szCs), which python-docx's
    # high-level API does NOT update, to keep it consistent with sz. ----
    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.append(szCs)
    szCs.set(qn("w:val"), str(size_pt * 2))

    # ---- Color: force to "auto" (Word's "Automatic" color, renders as
    # black), removing any themeColor/themeTint/themeShade references. ----
    color = rPr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rPr.append(color)
    for theme_attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
        if color.get(qn(theme_attr)) is not None:
            del color.attrib[qn(theme_attr)]
    color.set(qn("w:val"), "auto")


def _apply_base_styles(document: Document) -> None:
    """Shared styling setup used by ALL output formats: fixes the OOXML
    <w:zoom> validation issue, and forces Times New Roman 12pt, black
    (Automatic) font color across the Normal style AND the Heading 1/2
    styles, so ALL text in the document (body text and headings alike)
    uses the same font/size/color - overriding Word's theme-based heading
    defaults (see _hard_set_style_font for why this requires more than
    just setting .font.name/.font.size)."""
    # python-docx's default template omits the required w:percent attribute
    # on <w:zoom>; add it so the resulting file passes strict OOXML validation.
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")

    _hard_set_style_font(document.styles["Normal"], FONT_NAME, FONT_SIZE)
    style = document.styles["Normal"]
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)

    for heading_style_id in ("Heading 1", "Heading 2"):
        if heading_style_id in document.styles:
            _hard_set_style_font(document.styles[heading_style_id], FONT_NAME, FONT_SIZE)
